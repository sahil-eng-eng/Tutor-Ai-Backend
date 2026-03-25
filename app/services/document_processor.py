"""
Document Processor — extracts text from uploaded PDFs, DOCX, and plain text files.
"""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger("ai_tutor")


async def extract_text_from_file(file_path: str, file_type: str) -> str:
    """
    Extract text content from a file based on its type.
    Supports: PDF, DOCX, TXT.
    """
    extractors = {
        "pdf": _extract_pdf,
        "document": _extract_docx,
        "text": _extract_text,
        "image": _extract_image_ocr,
    }
    extractor = extractors.get(file_type, _extract_text)
    return await extractor(file_path)


async def _extract_pdf(file_path: str) -> str:
    """Extract text from a PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF

        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"PDF not found: {file_path}")

        doc = fitz.open(str(path))
        pages = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text.strip():
                pages.append(f"--- Page {page_num + 1} ---\n{text.strip()}")
        doc.close()

        if not pages:
            return "[PDF contained no extractable text — may be image-based]"
        return "\n\n".join(pages)
    except ImportError:
        logger.warning("PyMuPDF (fitz) not installed — PDF extraction unavailable")
        return "[PDF extraction unavailable — PyMuPDF not installed]"
    except Exception:
        logger.exception("PDF extraction failed for %s", file_path)
        return "[PDF extraction failed]"


async def _extract_docx(file_path: str) -> str:
    """Extract text from a DOCX file using python-docx."""
    try:
        from docx import Document

        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"DOCX not found: {file_path}")

        doc = Document(str(path))
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # Also extract table content
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))

        if not paragraphs:
            return "[DOCX contained no extractable text]"
        return "\n\n".join(paragraphs)
    except ImportError:
        logger.warning("python-docx not installed — DOCX extraction unavailable")
        return "[DOCX extraction unavailable — python-docx not installed]"
    except Exception:
        logger.exception("DOCX extraction failed for %s", file_path)
        return "[DOCX extraction failed]"


async def _extract_text(file_path: str) -> str:
    """Read plain text file."""
    try:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        content = path.read_text(encoding="utf-8", errors="replace")
        return content.strip() if content.strip() else "[File contained no text]"
    except Exception:
        logger.exception("Text extraction failed for %s", file_path)
        return "[Text extraction failed]"


async def _extract_image_ocr(file_path: str) -> str:
    """Extract text from an image using OCR (delegates to ocr_service)."""
    from app.services.ocr_service import extract_text_from_image

    return await extract_text_from_image(file_path)
